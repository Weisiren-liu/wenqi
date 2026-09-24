"""P0 + P1 行为测试：配置路径收敛、预设注册表统一（schema 2 → 3）。

对应 docs/ARCHITECTURE.md 3.1 / 3.2 与 docs/HANDOFF.md 第 4 节的验收标准：
- 两侧（GUI / formatter）读到同一个配置文件
- 内置预设加载时注入注册表、持久化时过滤
- active_preset_id 允许指向内置 id
- schema 2 → 3 迁移不丢用户预设、默认选中不变
- format_document() 统一接收完整 preset dict，旧调用方式保持兼容
"""

import json
from copy import deepcopy

from docx import Document

import docformat_gui
from scripts import formatter
from scripts.formatter import PRESETS, format_document


def _make_user_preset(preset_id, name, **overrides):
    preset = deepcopy(docformat_gui.DEFAULT_CUSTOM_SETTINGS)
    preset.update({"id": preset_id, "name": name})
    preset.update(overrides)
    return preset


def _make_docx(path):
    doc = Document()
    doc.add_paragraph("这是正文内容。")
    doc.save(path)


# ---------- P0：配置路径收敛 ----------

def test_gui_and_formatter_read_the_same_config_file(tmp_path, monkeypatch):
    appdata = tmp_path / "appdata"
    monkeypatch.setenv("APPDATA", str(appdata))
    monkeypatch.setattr(formatter.sys, "platform", "win32")

    gui_dir = docformat_gui._get_config_dir()
    assert gui_dir == formatter.get_config_dir()
    assert gui_dir == appdata / "DocFormatter"
    assert formatter.get_config_file() == gui_dir / "custom_settings.json"

    # GUI 写、formatter 读 —— 必须命中同一个文件
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", formatter.get_config_file())
    preset = _make_user_preset("shared-1", "两侧共享")
    preset["page"] = {"top": 1.1, "bottom": 1.2, "left": 1.3, "right": 1.4}
    docformat_gui.save_custom_settings({
        "schema_version": docformat_gui.CONFIG_SCHEMA_VERSION,
        "active_preset_id": "shared-1",
        "presets": [preset],
    })

    loaded_by_formatter = formatter.load_custom_preset()
    assert loaded_by_formatter is not None
    assert loaded_by_formatter["id"] == "shared-1"
    assert loaded_by_formatter["page"]["left"] == 1.3


# ---------- P1：内置预设注入注册表 ----------

def test_load_injects_builtin_presets_as_readonly_entries(tmp_path, monkeypatch):
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", tmp_path / "custom_settings.json")

    config = docformat_gui.load_custom_settings()
    ids = [p.get("id") for p in config["presets"]]
    assert ids[:3] == ["official", "academic", "legal"]

    for entry in config["presets"][:3]:
        assert entry["is_builtin"] is True
    # 注入的内置条目与 formatter.PRESETS 定义逐字一致（不合并默认值，不改变格式化行为）
    for builtin_id in ("official", "academic", "legal"):
        entry = docformat_gui.get_preset_by_id(config, builtin_id)
        definition_keys = set(PRESETS[builtin_id]) - {"id", "is_builtin"}
        for key in definition_keys:
            assert entry[key] == PRESETS[builtin_id][key]

    # 全新安装：注册表 = 3 内置 + 1 用户，且默认选中的仍是用户项（自定义格式）
    user = [p for p in config["presets"] if not p["is_builtin"]]
    assert len(user) == 1
    assert config["active_preset_id"] == user[0]["id"]


def test_save_strips_builtin_entries_from_user_config(tmp_path, monkeypatch):
    config_file = tmp_path / "custom_settings.json"
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", config_file)

    config = docformat_gui.load_custom_settings()
    docformat_gui.save_custom_settings(config)

    on_disk = json.loads(config_file.read_text(encoding="utf-8"))
    assert on_disk["schema_version"] == 3
    assert not any(p.get("is_builtin") for p in on_disk["presets"])
    assert not any(p.get("id") in ("official", "academic", "legal")
                   for p in on_disk["presets"])
    assert "builtin_overrides" in on_disk

    # 保存不改调用方的内存注册表
    assert docformat_gui.get_preset_by_id(config, "official") is not None


def test_active_preset_id_may_point_to_builtin_id(tmp_path, monkeypatch):
    config_file = tmp_path / "custom_settings.json"
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", config_file)

    docformat_gui.save_custom_settings({
        "schema_version": docformat_gui.CONFIG_SCHEMA_VERSION,
        "active_preset_id": "official",
        "presets": [_make_user_preset("user-1", "我的格式")],
    })

    loaded = docformat_gui.load_custom_settings()
    # 不再被 load_custom_settings 校验重置
    assert loaded["active_preset_id"] == "official"
    active = docformat_gui.get_active_preset(loaded)
    assert active["is_builtin"] is True
    assert active["body"] == PRESETS["official"]["body"]

    # 存盘再读，仍指向内置 id
    docformat_gui.save_custom_settings(loaded)
    reloaded = docformat_gui.load_custom_settings()
    assert reloaded["active_preset_id"] == "official"


def test_user_preset_wins_on_builtin_id_collision(tmp_path, monkeypatch):
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", tmp_path / "custom_settings.json")

    docformat_gui.save_custom_settings({
        "schema_version": docformat_gui.CONFIG_SCHEMA_VERSION,
        "active_preset_id": "official",
        "presets": [_make_user_preset("official", "我的官方格式")],
    })

    loaded = docformat_gui.load_custom_settings()
    official_entries = [p for p in loaded["presets"] if p.get("id") == "official"]
    assert len(official_entries) == 1
    assert official_entries[0]["name"] == "我的官方格式"
    assert official_entries[0]["is_builtin"] is False
    # 其余内置项照常注入
    assert any(p["id"] == "academic" and p["is_builtin"] for p in loaded["presets"])


# ---------- P1：schema 2 → 3 迁移 ----------

def test_schema2_config_migrates_without_losing_user_presets(tmp_path, monkeypatch):
    config_file = tmp_path / "custom_settings.json"
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", config_file)

    schema2 = {
        "schema_version": 2,
        "active_preset_id": "preset-b",
        "page_number_config_version": 1,
        "presets": [
            _make_user_preset("preset-a", "A"),
            _make_user_preset("preset-b", "B"),
        ],
    }
    schema2["presets"][0]["page"]["left"] = 1.1
    schema2["presets"][1]["page"]["left"] = 2.2
    config_file.write_text(json.dumps(schema2, ensure_ascii=False), encoding="utf-8")

    loaded = docformat_gui.load_custom_settings()
    assert loaded["schema_version"] == 3

    user = {p["id"]: p for p in loaded["presets"] if not p["is_builtin"]}
    assert set(user) == {"preset-a", "preset-b"}
    assert user["preset-a"]["name"] == "A"
    assert user["preset-b"]["page"]["left"] == 2.2
    # 默认选中格式不变
    assert loaded["active_preset_id"] == "preset-b"
    assert docformat_gui.get_active_preset(loaded)["name"] == "B"

    # 升级后可正常保存，用户预设全保留
    docformat_gui.save_custom_settings(loaded)
    on_disk = json.loads(config_file.read_text(encoding="utf-8"))
    assert on_disk["schema_version"] == 3
    assert [p["id"] for p in on_disk["presets"]] == ["preset-a", "preset-b"]


def test_formatter_load_custom_preset_accepts_schema3(tmp_path, monkeypatch):
    appdata = tmp_path / "appdata"
    monkeypatch.setenv("APPDATA", str(appdata))
    monkeypatch.setattr(formatter.sys, "platform", "win32")
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", formatter.get_config_file())

    docformat_gui.save_custom_settings({
        "schema_version": docformat_gui.CONFIG_SCHEMA_VERSION,
        "active_preset_id": "user-x",
        "presets": [
            _make_user_preset("user-other", "另一个"),
            _make_user_preset("user-x", "选中的"),
        ],
    })
    loaded = formatter.load_custom_preset()
    assert loaded is not None
    assert loaded["id"] == "user-x"
    assert loaded["name"] == "选中的"


# ---------- P1：format_document 统一接收完整 preset dict ----------

def test_format_document_accepts_full_preset_dict(tmp_path):
    preset = deepcopy(PRESETS["official"])
    preset.update({"id": "reg-entry", "is_builtin": False, "name": "注册表条目"})
    preset["page"] = {"top": 1.1, "bottom": 1.2, "left": 1.3, "right": 1.4}

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_docx(source)

    # 同时给 preset_name，验证 preset dict 优先
    format_document(str(source), str(output),
                    preset_name="official", preset=preset)

    section = Document(output).sections[0]
    assert abs(section.top_margin.cm - 1.1) < 0.02
    assert abs(section.bottom_margin.cm - 1.2) < 0.02
    assert abs(section.left_margin.cm - 1.3) < 0.02
    assert abs(section.right_margin.cm - 1.4) < 0.02


def test_builtin_registry_entry_formats_identically_to_preset_name(tmp_path, monkeypatch):
    monkeypatch.setattr(docformat_gui, "CONFIG_FILE", tmp_path / "custom_settings.json")
    config = docformat_gui.load_custom_settings()
    official = docformat_gui.get_preset_by_id(config, "official")

    source = tmp_path / "source.docx"
    out_by_name = tmp_path / "by_name.docx"
    out_by_dict = tmp_path / "by_dict.docx"
    _make_docx(source)

    format_document(str(source), str(out_by_name), preset_name="official")
    format_document(str(source), str(out_by_dict), preset=official)

    a = Document(out_by_name)
    b = Document(out_by_dict)
    for x, y in zip(a.sections, b.sections):
        assert x.top_margin.cm == y.top_margin.cm
        assert x.bottom_margin.cm == y.bottom_margin.cm
        assert x.left_margin.cm == y.left_margin.cm
        assert x.right_margin.cm == y.right_margin.cm
    for x, y in zip(a.paragraphs, b.paragraphs):
        assert x.text == y.text
        assert x.alignment == y.alignment
        for rx, ry in zip(x.runs, y.runs):
            assert rx.font.name == ry.font.name
            assert rx.font.size == ry.font.size
            assert rx.bold == ry.bold


def test_legacy_preset_name_calls_still_work(tmp_path, monkeypatch):
    """旧调用方式（preset_name='custom' + custom_settings）行为保持不变。"""
    custom = deepcopy(PRESETS["official"])
    custom["page"] = {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0}

    source = tmp_path / "source.docx"
    output = tmp_path / "output.docx"
    _make_docx(source)

    format_document(str(source), str(output),
                    preset_name="custom", custom_settings=custom)
    section = Document(output).sections[0]
    assert abs(section.top_margin.cm - 2.0) < 0.02
