"""Verify that the new desktop shell preserves actual document workflows."""

from copy import deepcopy
from pathlib import Path
import time

from docx import Document
import pytest

import docformat_gui as model
from scripts import desktop_ui


@pytest.fixture
def api(tmp_path, monkeypatch):
    monkeypatch.setattr(model, "CONFIG_FILE", tmp_path / "settings" / "custom_settings.json")
    monkeypatch.setattr(desktop_ui, "_fonts", lambda: ["仿宋", "宋体", "Arial"])
    return desktop_ui.DesktopApi()


def document(path, text="这是正文, 请做好工作!"):
    doc = Document()
    doc.add_paragraph("工作安排")
    doc.add_paragraph(text)
    doc.save(path)
    return str(path)


def wait_for_result(api):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        state = api.get_state()
        if not state["job"]["running"]:
            return state
        time.sleep(.02)
    pytest.fail("Document worker did not finish")


def test_edits_persist_and_restore_without_changing_the_base_preset(api):
    original = deepcopy(model.DEFAULT_CUSTOM_SETTINGS["body"])
    api.update_setting("body", "size", 18)
    api.update_setting("body", "font_en", "Arial")
    api.update_setting("body", "line_spacing", "1.5倍")
    restarted = desktop_ui.DesktopApi()
    assert restarted.get_state()["effective"]["body"]["size"] == 18
    assert restarted.get_state()["effective"]["body"]["font_en"] == "Arial"
    assert restarted.get_state()["effective"]["body"]["line_spacing"] == "1.5倍"
    restarted.update_setting("body", "line_spacing", "")
    assert restarted.get_state()["effective"]["body"]["line_spacing"] == original["line_spacing"]
    assert model.DEFAULT_CUSTOM_SETTINGS["body"] == original
    restarted.restore_settings("body")
    assert restarted.get_state()["effective"]["body"] == original


def test_custom_format_snapshot_and_switching(api):
    api.update_setting("page", "left", 2.1)
    result = api.save_preset("我的信函")
    custom_id = result["active"]
    assert custom_id != "custom"
    api.select_preset("official")
    api.select_preset(custom_id)
    assert api.get_state()["effective"]["page"]["left"] == 2.1
    api.manage_preset("rename", "信函格式")
    assert api.get_state()["effective"]["name"] == "信函格式"
    api.manage_preset("delete")
    assert api.get_state()["active"] == "custom"


@pytest.mark.parametrize("mode", ["smart", "punctuation", "analyze"])
def test_all_three_processing_modes_preserve_the_original(api, tmp_path, mode):
    source = document(tmp_path / "原文 有空格.docx")
    original = Path(source).read_bytes()
    api.add_files([source])
    assert api.get_state()["preview"][0]["text"] == "工作安排"
    api.start_processing({"mode": mode})
    result = wait_for_result(api)["results"][0]
    assert result["status"] == "success", result
    assert Path(source).read_bytes() == original
    if mode == "analyze":
        assert set(result["diagnosis"]) == {"punctuation", "numbering", "paragraph", "font"}
        assert list(tmp_path.glob("*_processed.docx")) == []
    else:
        output = Document(result["path"])
        assert "这是正文，" in output.paragraphs[1].text
        assert output.paragraphs[1].text.endswith("！")


def test_actual_output_uses_the_effective_custom_parameters(api, tmp_path):
    source = document(tmp_path / "source.docx")
    api.add_files([source])
    api.update_setting("body", "size", 18)
    api.update_setting("body", "font_en", "Arial")
    api.update_setting("page", "left", 2.1)
    api.start_processing({"mode": "smart"})
    result = wait_for_result(api)["results"][0]
    assert result["status"] == "success", result
    output = Document(result["path"])
    assert abs(output.sections[0].left_margin.cm - 2.1) < .02
    assert output.paragraphs[1].runs[0].font.size.pt == 18
    assert output.paragraphs[1].runs[0].font.name == "Arial"


def test_batch_avoids_overwriting_existing_output(api, tmp_path):
    first = document(tmp_path / "a.docx")
    second = document(tmp_path / "b.docx")
    existing = tmp_path / "a_processed.docx"
    existing.write_bytes(b"keep this existing output")
    api.add_files([first, second, first])
    assert len(api.get_state()["files"]) == 2
    api.start_processing({"mode": "punctuation"})
    state = wait_for_result(api)
    assert len(state["results"]) == 2
    assert all(result["status"] == "success" for result in state["results"])
    assert existing.read_bytes() == b"keep this existing output"
    assert Path(state["results"][0]["path"]).name == "a_processed_2.docx"


def test_invalid_document_reports_error_and_releases_the_worker(api, tmp_path):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a docx")
    api.add_files([str(source)])
    api.start_processing({"mode": "smart"})
    state = wait_for_result(api)
    assert state["results"][0]["status"] == "error"
    assert state["results"][0]["error"]
    assert state["job"]["running"] is False
    assert not list(tmp_path.glob("*_processed.docx"))


def test_unsafe_or_invalid_settings_are_rejected(api):
    for arguments in [("body", "size", -1), ("body", "size", 0),
                      ("body", "align", "invalid"), ("page", "top", float("nan")),
                      ("page", "left", 25), ("body", "line_spacing", "20倍")]:
        with pytest.raises(ValueError):
            api.update_setting(*arguments)
    with pytest.raises(ValueError):
        api.start_processing({"mode": "smart"})


def test_original_input_cannot_be_used_as_output(api, tmp_path):
    source = document(tmp_path / "source.docx")
    api.add_files([source])
    api._output = {"kind": "file", "path": source}
    with pytest.raises(ValueError, match="原文件"):
        api.start_processing({"mode": "smart"})


def test_pasted_document_becomes_a_real_input_and_temp_file_is_cleaned(api, tmp_path, monkeypatch):
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    state = api.create_from_text("文本标题", "第一段\n\n第二段", False)
    source = Path(state["files"][0]["path"])
    assert source.exists()
    assert Document(source).paragraphs[0].text == "文本标题"
    api.start_processing({"mode": "smart"})
    result = wait_for_result(api)["results"][0]
    assert result["status"] == "success", result
    assert Path(result["path"]).parent == tmp_path / "Documents"
    api._cleanup()
    assert not source.exists()
    assert Path(result["path"]).exists()


def test_cancellation_finishes_current_file_and_skips_the_remaining_queue(api, tmp_path, monkeypatch):
    files = [document(tmp_path / f"{index}.docx") for index in range(3)]
    api.add_files(files)
    original = api._process_one

    def process_and_cancel(*args):
        result = original(*args)
        api.cancel_processing()
        return result

    monkeypatch.setattr(api, "_process_one", process_and_cancel)
    api.start_processing({"mode": "punctuation"})
    state = wait_for_result(api)
    assert state["job"]["cancelled"] is True
    assert len(state["results"]) == 1
    assert not (tmp_path / "1_processed.docx").exists()
