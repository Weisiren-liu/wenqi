"""Local desktop shell and document workflows for the Liquid Glass interface."""

from copy import deepcopy
from datetime import datetime
import json
import logging
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import threading

from docx import Document

import docformat_gui as model
from scripts.branding import APP_NAME, VERSION
from scripts.formatter import format_document, detect_para_type


ELEMENTS = [
    ("title", "主标题"), ("recipient", "主送机关"),
    ("heading1", "一级标题"), ("heading2", "二级标题"),
    ("heading3", "三级标题"), ("heading4", "四级标题"),
    ("body", "正文"), ("signature", "落款"), ("date", "日期"),
    ("attachment", "附件"), ("closing", "结束语"),
]
SAMPLE = [
    {"text": "关于做好年度工作安排的通知", "element": "title"},
    {"text": "各部门：", "element": "recipient"},
    {"text": "为有序推进各项工作，现将有关事项通知如下。", "element": "body"},
    {"text": "一、明确工作目标", "element": "heading1"},
    {"text": "结合实际情况，细化工作任务，明确时间安排与职责分工，确保各项工作落到实处。", "element": "body"},
    {"text": "二、做好统筹协调", "element": "heading1"},
    {"text": "加强沟通协作，及时总结进展，持续提升工作质量。", "element": "body"},
    {"text": "办公室", "element": "signature"},
    {"text": "2026年10月12日", "element": "date"},
]


def _resource(relative):
    return Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1])) / relative


def _fonts():
    names = set(model._PRIORITY_FONTS)
    if sys.platform == "win32":
        import winreg
        try:
            with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE,
                               r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts") as key:
                for index in range(winreg.QueryInfoKey(key)[1]):
                    name = winreg.EnumValue(key, index)[0].split(" (")[0]
                    names.update(name.split(" & "))
        except OSError:
            pass
    else:
        try:
            result = subprocess.run(["fc-list", ":", "family"], capture_output=True,
                                    text=True, timeout=3)
            names.update(name.strip() for line in result.stdout.splitlines()
                         for name in line.split(",") if name.strip())
        except (OSError, subprocess.TimeoutExpired):
            pass
    return sorted(names, key=str.casefold)


class DesktopApi:
    """All exposed actions operate on local files, through the existing preset model."""

    def __init__(self):
        self._lock = threading.RLock()
        self._config = model.load_custom_settings()
        self._files = []
        self._preview = deepcopy(SAMPLE)
        self._output = {"kind": "beside", "path": ""}
        self._logs = []
        self._results = []
        self._job = {"running": False, "percent": 0, "stage": "", "cancelled": False}
        self._cancel = threading.Event()
        self._window = None
        self._window_controls = None
        self._window_maximized = False
        self._window_minimized = False
        self._window_closing = False
        self._fonts = _fonts()
        self._temporary_inputs = []

    def _attach(self, window):
        self._window = window

    def _idle(self):
        if self._job["running"]:
            raise ValueError("请等待当前处理完成。")

    def _persist(self):
        model.CONFIG_FILE.parent.mkdir(parents=True, exist_ok=True)
        if model.save_custom_settings(self._config) is None:
            raise OSError("无法保存格式设置，请检查配置目录的写入权限。")

    def _log(self, message, level="info"):
        with self._lock:
            self._logs.append({"time": datetime.now().strftime("%H:%M:%S"),
                               "message": str(message), "level": level})
            self._logs = self._logs[-1000:]

    def get_state(self):
        with self._lock:
            active = self._config["active_preset_id"]
            return {
                "version": VERSION,
                "files": [{"path": path, "name": Path(path).name,
                           "size": Path(path).stat().st_size if Path(path).exists() else 0}
                          for path in self._files],
                "presets": [{"id": p["id"], "name": p["name"],
                             "builtin": bool(p.get("is_builtin"))}
                            for p in self._config["presets"]],
                "active": active,
                "effective": model.get_effective_preset(self._config, active),
                "overrides": deepcopy(model.get_preset_overrides(self._config, active)),
                "fonts": self._fonts,
                "elements": ELEMENTS,
                "preview": self._preview,
                "previewIsSample": not self._files or Path(self._files[0]).suffix.lower() != ".docx",
                "output": dict(self._output),
                "job": dict(self._job),
                "results": deepcopy(self._results),
                "logs": list(self._logs),
                "canConvert": sys.platform == "win32",
            }

    def _read_preview(self):
        self._preview = deepcopy(SAMPLE)
        if self._files and Path(self._files[0]).suffix.lower() == ".docx":
            try:
                doc = Document(self._files[0])
                self._preview = []
                paragraphs = [p for p in doc.paragraphs if p.text.strip()]
                texts = [p.text for p in paragraphs]
                previous = None
                for index, paragraph in enumerate(paragraphs):
                    kind = detect_para_type(paragraph.text, index, len(paragraphs),
                                            paragraph.alignment, texts, prev_para_type=previous)
                    previous = kind
                    self._preview.append({"text": paragraph.text[:500], "element": kind})
                    if len(self._preview) >= 22:
                        break
                if not self._preview:
                    self._preview = [{"text": "此文档没有可预览的正文段落。", "element": "body"}]
            except Exception:
                self._preview = [{"text": "无法读取预览，请检查文件是否有效。", "element": "body"}]

    def add_files(self, paths):
        with self._lock:
            self._idle()
            accepted = []
            for raw in paths:
                path = Path(raw).expanduser().resolve()
                if path.is_file() and path.suffix.lower() in (".docx", ".doc", ".wps") and not path.name.startswith("~$"):
                    if path.suffix.lower() != ".docx" and sys.platform != "win32":
                        raise ValueError("请先将 .doc / .wps 文件转换为 .docx。")
                    if str(path) not in self._files:
                        accepted.append(str(path))
            if not accepted and not self._files:
                raise ValueError("请选择有效的 Word 文档（.docx、.doc 或 .wps）。")
            self._files.extend(dict.fromkeys(accepted))
            if len(self._files) > 1 and self._output["kind"] == "file":
                self._output = {"kind": "beside", "path": ""}
            self._results = []
            self._job.update(percent=0, stage="", cancelled=False)
            self._read_preview()
            return self.get_state()

    def choose_files(self):
        import webview
        with self._lock:
            self._idle()
        paths = self._window.create_file_dialog(webview.FileDialog.OPEN, allow_multiple=True,
                                               file_types=("Word 文档 (*.docx;*.doc;*.wps)",))
        return self.add_files(paths) if paths else self.get_state()

    def choose_folder(self):
        import webview
        with self._lock:
            self._idle()
        paths = self._window.create_file_dialog(webview.FileDialog.FOLDER)
        if paths:
            files = [str(p) for p in sorted(Path(paths[0]).rglob("*"))
                     if p.is_file() and p.suffix.lower() in (".docx", ".doc", ".wps") and not p.name.startswith("~$")]
            if not files:
                raise ValueError("这个文件夹中没有 Word 文档。")
            return self.add_files(files)
        return self.get_state()

    def remove_file(self, path):
        with self._lock:
            self._idle()
            if path in self._files:
                self._files.remove(path)
            self._results = []
            self._read_preview()
            return self.get_state()

    def clear_files(self):
        with self._lock:
            self._idle()
            self._files.clear()
            self._results.clear()
            self._preview = deepcopy(SAMPLE)
            return self.get_state()

    def choose_output(self, kind="folder"):
        import webview
        with self._lock:
            self._idle()
        if kind == "beside":
            self._output = {"kind": "beside", "path": ""}
        elif kind == "file" and len(self._files) == 1:
            first = Path(self._files[0])
            paths = self._window.create_file_dialog(webview.FileDialog.SAVE, directory=str(first.parent),
                save_filename=f"{first.stem}_processed{first.suffix}",
                file_types=("Word 文档 (*.docx;*.doc;*.wps)",))
            if paths:
                selected = Path(paths[0]).resolve()
                if selected.suffix.lower() not in (".docx", ".doc", ".wps"):
                    raise ValueError("输出文件需要使用 .docx、.doc 或 .wps 扩展名。")
                if selected == first.resolve():
                    raise ValueError("输出文件不能与原文件相同。")
                self._output = {"kind": "file", "path": str(selected)}
        else:
            paths = self._window.create_file_dialog(webview.FileDialog.FOLDER)
            if paths:
                self._output = {"kind": "folder", "path": str(Path(paths[0]).resolve())}
        return self.get_state()

    def select_preset(self, preset_id):
        with self._lock:
            self._idle()
            if model.get_preset_by_id(self._config, preset_id) is None:
                raise ValueError("找不到这个格式。")
            self._config["active_preset_id"] = preset_id
            self._persist()
            return self.get_state()

    def update_setting(self, key, field, value):
        with self._lock:
            self._idle()
            defaults = model.DEFAULT_CUSTOM_SETTINGS
            if key not in defaults or key == "name":
                raise ValueError("未知格式参数。")
            default = defaults[key]
            if isinstance(default, dict):
                if field not in default:
                    raise ValueError("未知格式参数。")
                default = default[field]
            elif field is not None:
                raise ValueError("这个参数没有子项。")
            if isinstance(value, str) and not value.strip():
                value = None
            if value is not None:
                if isinstance(default, bool):
                    if not isinstance(value, bool):
                        raise ValueError("开关参数需要布尔值。")
                elif field == "line_spacing" and isinstance(value, str) and value.strip().endswith("倍"):
                    multiple = float(value.strip()[:-1])
                    if not math.isfinite(multiple) or not 0 < multiple <= 10:
                        raise ValueError("多倍行距需要大于零且不超过 10 倍。")
                    value = f"{multiple:g}倍"
                elif isinstance(default, (int, float)):
                    value = float(value)
                    if not math.isfinite(value) or value < 0 or value > 576:
                        raise ValueError("请输入 0 到 576 之间的数值。")
                    if (field == "size" or key == "page_number_size") and value == 0:
                        raise ValueError("字号需要大于零。")
                    if key == "page" and value > 20:
                        raise ValueError("页边距不能超过 20 厘米。")
                    if key == "page_number_offset_mm" and value > 30:
                        raise ValueError("页码偏移不能超过 30 毫米。")
                else:
                    value = str(value).strip()
                    if not value:
                        value = None
                    elif len(value) > 100:
                        raise ValueError("参数内容过长。")
                choices = {
                    "align": {"left", "center", "right", "justify"},
                    "space_handling": {"remove_all", "keep_all", "keep_en_boundary"},
                    "page_number_style": set(model.PAGE_NUMBER_STYLE_OPTIONS),
                    "page_number_position": set(model.PAGE_NUMBER_POSITION_OPTIONS),
                }
                option = field if field == "align" else key
                if value is not None and option in choices and value not in choices[option]:
                    raise ValueError("请选择有效选项。")
            before = deepcopy(self._config)
            try:
                model.set_preset_override(self._config, self._config["active_preset_id"], key, field, value)
                self._persist()
            except Exception:
                self._config = before
                raise
            return self.get_state()

    def restore_settings(self, element=None):
        with self._lock:
            self._idle()
            active = self._config["active_preset_id"]
            if element:
                model.clear_preset_override_row(self._config, active, element)
            else:
                model.clear_preset_overrides(self._config, active)
            self._persist()
            return self.get_state()

    def save_preset(self, name):
        with self._lock:
            self._idle()
            name = str(name).strip()
            if not name or len(name) > 60:
                raise ValueError("格式名称需要 1 到 60 个字符。")
            if any(p["name"] == name for p in self._config["presets"]):
                raise ValueError("已有同名格式，请换一个名称。")
            model.save_as_new_preset(self._config, self._config["active_preset_id"], name)
            self._persist()
            return self.get_state()

    def manage_preset(self, action, name=""):
        import webview
        with self._lock:
            self._idle()
            active = self._config["active_preset_id"]
            preset = model.get_preset_by_id(self._config, active)
            if action in ("rename", "delete") and preset.get("is_builtin"):
                raise ValueError("起始格式不能重命名或删除，请先存为新格式。")
            if action == "rename":
                name = name.strip()
                if not name or len(name) > 60 or any(p["name"] == name and p["id"] != active for p in self._config["presets"]):
                    raise ValueError("请使用有效且不重复的名称。")
                preset["name"] = name
            elif action == "delete":
                self._config["presets"] = [p for p in self._config["presets"] if p["id"] != active]
                self._config["preset_overrides"].pop(active, None)
                self._config["active_preset_id"] = "custom"
            elif action == "export":
                paths = self._window.create_file_dialog(webview.FileDialog.SAVE,
                    save_filename=f"{preset['name']}.json", file_types=("格式文件 (*.json)",))
                if paths:
                    Path(paths[0]).write_text(json.dumps(model.get_effective_preset(self._config, active),
                                                       ensure_ascii=False, indent=2), encoding="utf-8")
            elif action == "import":
                paths = self._window.create_file_dialog(webview.FileDialog.OPEN, file_types=("格式文件 (*.json)",))
                if paths:
                    data = json.loads(Path(paths[0]).read_text(encoding="utf-8"))
                    if isinstance(data, dict) and "presets" in data:
                        entries = data["presets"]
                    else:
                        entries = [data]
                    if not isinstance(entries, list) or not entries or any(not isinstance(p, dict) for p in entries):
                        raise ValueError("这个文件不是有效的格式文件。")
                    import uuid
                    for entry in entries:
                        if not any(key in entry for key, _ in ELEMENTS):
                            raise ValueError("这个文件没有段落格式参数。")
                        merged = model._merge_settings(model.DEFAULT_CUSTOM_SETTINGS, entry)
                        merged.update(id=str(uuid.uuid4()), is_builtin=False)
                        merged["name"] = str(entry.get("name", "导入格式"))[:60]
                        self._config["presets"].append(merged)
                        self._config["active_preset_id"] = merged["id"]
            else:
                raise ValueError("未知格式操作。")
            self._persist()
            return self.get_state()

    def create_from_text(self, title, body, markdown=False):
        with self._lock:
            self._idle()
            if not str(body).strip():
                raise ValueError("请先输入正文。")
            safe_title = str(title).strip() or "未命名文档"
            folder = Path(tempfile.mkdtemp(prefix="wenqi-paste-"))
            safe_name = "".join(c for c in safe_title if c not in '<>:"/\\|?*')[:60] or "未命名文档"
            path = folder / f"{safe_name}.docx"
            creator = model._create_docx_from_markdown if markdown else model._create_docx_from_text
            creator(str(title).strip(), str(body), str(path))
            self._temporary_inputs.append(str(path))
            if self._output["kind"] == "beside":
                self._output = {"kind": "folder", "path": str(Path.home() / "Documents")}
            return self.add_files([str(path)])

    def start_processing(self, options):
        with self._lock:
            self._idle()
            if not self._files:
                raise ValueError("请先添加文档。")
            mode = options.get("mode", "smart")
            if mode not in ("smart", "punctuation", "analyze"):
                raise ValueError("未知处理方式。")
            if any(not Path(p).is_file() for p in self._files):
                raise ValueError("有文件已被移动或删除，请重新添加。")
            preset = deepcopy(model.get_effective_preset(self._config, self._config["active_preset_id"]))
            files = list(self._files)
            output = dict(self._output)
            if output["kind"] == "file" and Path(output["path"]).resolve() in {Path(p).resolve() for p in files}:
                raise ValueError("输出文件不能与原文件相同。")
            self._cancel.clear()
            self._results = []
            self._logs = []
            self._job.update(running=True, percent=0, stage="准备处理", cancelled=False)
            thread = threading.Thread(target=self._process_batch,
                args=(files, output, mode, preset, bool(options.get("revision")), bool(options.get("autoOpen"))),
                daemon=True, name="wenqi-document-worker")
            thread.start()
            return self.get_state()

    def cancel_processing(self):
        self._cancel.set()
        with self._lock:
            if self._job["running"]:
                self._job["stage"] = "当前文档完成后停止"
            return self.get_state()

    def _progress(self, percent, stage):
        with self._lock:
            self._job["percent"] = max(0, min(100, round(percent)))
            if not self._cancel.is_set():
                self._job["stage"] = str(stage)

    @staticmethod
    def _unique_output(path, source):
        path = Path(path)
        if path.resolve() == Path(source).resolve():
            raise ValueError("输出文件不能与原文件相同。")
        base = path
        number = 2
        while path.exists():
            path = base.with_name(f"{base.stem}_{number}{base.suffix}")
            number += 1
        return path

    def _process_batch(self, files, output, mode, preset, revision, auto_open):
        owner = self

        class Handler(logging.Handler):
            def emit(self, record):
                owner._log(record.getMessage())

        logger = logging.getLogger("docformat.formatter")
        handler = Handler()
        logger.addHandler(handler)
        logger.setLevel(logging.INFO)
        try:
            for index, source in enumerate(files):
                if self._cancel.is_set():
                    break
                self._log(f"开始处理：{Path(source).name}")
                if output["kind"] == "file":
                    destination = Path(output["path"])
                else:
                    directory = Path(output["path"]) if output["kind"] == "folder" else Path(source).parent
                    destination = directory / f"{Path(source).stem}_processed{Path(source).suffix}"
                try:
                    if mode != "analyze":
                        destination.parent.mkdir(parents=True, exist_ok=True)
                        destination = self._unique_output(destination, source)
                    progress = lambda percent, stage, offset=index: self._progress((offset + percent / 100) / len(files) * 100, stage)
                    result = self._process_one(source, destination, mode, preset, revision, progress)
                    result.update(name=Path(source).name, source=source, status="success")
                    self._log(f"{'诊断完成' if mode == 'analyze' else '已保存'}：{result.get('path', Path(source).name)}", "success")
                except Exception as exc:
                    result = {"name": Path(source).name, "source": source, "status": "error", "error": str(exc)}
                    self._log(f"{Path(source).name}：{exc}", "error")
                with self._lock:
                    self._results.append(result)
            with self._lock:
                self._job.update(running=False, cancelled=self._cancel.is_set(),
                    percent=self._job["percent"] if self._cancel.is_set() else 100,
                    stage="已停止" if self._cancel.is_set() else "处理完成")
            if auto_open and mode != "analyze" and len(self._results) == 1 and self._results[0].get("path"):
                self._open_local(self._results[0]["path"])
        finally:
            logger.removeHandler(handler)
            with self._lock:
                self._job["running"] = False

    def _process_one(self, source, destination, mode, preset, revision, progress):
        from scripts.punctuation import process_paragraph
        from scripts.east_asian_typography import apply_chinese_line_break_rules
        from scripts.analyzer import analyze_punctuation, analyze_numbering, analyze_paragraph_format, analyze_font
        converted = None
        try:
            with tempfile.TemporaryDirectory(prefix="wenqi-process-") as work:
                input_path = source
                if Path(source).suffix.lower() in (".doc", ".wps"):
                    from scripts.converter import convert_to_docx
                    progress(0, "转换为 Word 文档")
                    converted = convert_to_docx(source)
                    input_path = converted
                if mode == "analyze":
                    progress(20, "检查文档格式")
                    doc = Document(input_path)
                    issues = {"punctuation": analyze_punctuation(doc), "numbering": analyze_numbering(doc),
                              "paragraph": analyze_paragraph_format(doc), "font": analyze_font(doc)}
                    progress(100, "诊断完成")
                    return {"diagnosis": issues}
                doc = Document(input_path)
                paragraphs = list(doc.paragraphs)
                paragraphs += [p for table in doc.tables for row in table.rows for cell in row.cells for p in cell.paragraphs]
                changes = sum(bool(process_paragraph(p, space_mode=preset.get("space_handling", "remove_all"))) for p in paragraphs)
                apply_chinese_line_break_rules(doc)
                cleaned = str(Path(work) / "punctuation.docx")
                doc.save(cleaned)
                self._log(f"修复了 {changes} 段标点")
                progress(10, "应用文档格式" if mode == "smart" else "保存文档")
                needs_conversion = destination.suffix.lower() in (".doc", ".wps")
                target = Path(work) / "formatted.docx" if needs_conversion else destination
                if mode == "smart":
                    format_document(cleaned, str(target), preset=preset,
                        bold_serial=preset.get("bold_serial", True), revision_mode=revision,
                        progress_callback=lambda current, total, stage: progress(10 + current / max(1, total) * 80, stage))
                else:
                    import shutil
                    shutil.copyfile(cleaned, target)
                if needs_conversion:
                    from scripts.converter import convert_from_docx
                    progress(92, "保存原文档格式")
                    try:
                        actual = convert_from_docx(str(target), str(destination), format=destination.suffix[1:])
                        destination = Path(actual or destination)
                    except Exception as exc:
                        import shutil
                        fallback = self._unique_output(destination.with_suffix(".docx"), source)
                        shutil.copyfile(target, fallback)
                        destination = fallback
                        self._log(f"原格式保存失败，已保留 .docx：{exc}", "warning")
                progress(100, "完成")
                return {"path": str(destination), "changes": changes}
        finally:
            if converted and Path(converted).exists():
                Path(converted).unlink()

    @staticmethod
    def _open_local(path):
        if sys.platform == "win32":
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", path])
        else:
            subprocess.Popen(["xdg-open", path])

    def open_result(self, path, folder=False):
        with self._lock:
            permitted = {r.get("path") for r in self._results if r.get("path")}
            if path not in permitted:
                raise ValueError("找不到处理结果。")
        target = Path(path)
        if not target.exists():
            raise ValueError("这个文件已被移动或删除。")
        self._open_local(str(target.parent if folder else target))
        return True

    def window_action(self, action):
        if action == "close":
            if self._job["running"]:
                raise ValueError("请先停止处理，再关闭窗口。")
            self._window.destroy()
        elif action == "minimize":
            if self._window_controls:
                self._window_controls.action("minimize")
            else:
                self._window.minimize()
        elif action == "maximize":
            command = "restore" if self.window_state()["maximized"] else "maximize"
            if self._window_controls:
                self._window_controls.action(command)
            else:
                getattr(self._window, command)()
        else:
            raise ValueError("未知窗口操作。")
        return self.window_state()

    def window_state(self):
        if self._window_controls:
            return self._window_controls.state()
        return {"maximized": self._window_maximized, "minimized": self._window_minimized,
                "nativeDrag": False}

    def begin_window_drag(self):
        if self._window_controls:
            self._window_controls.begin_drag()
        return True

    def begin_window_resize(self, edge):
        if self._window_controls:
            self._window_controls.begin_resize(edge)
        return True

    def _publish_window_state(self, maximized=None, minimized=None):
        if self._window_closing:
            return
        if maximized is not None:
            self._window_maximized = maximized
        if minimized is not None:
            self._window_minimized = minimized
        if self._window.events.loaded.is_set() and not self.window_state()["minimized"]:
            try:
                self._window.evaluate_js("window.pywebview?.api.window_state().then(value => window.wenqi?.receiveWindowState(value))")
            except Exception:
                pass

    def _cleanup(self):
        for path in self._temporary_inputs:
            file = Path(path)
            if file.exists():
                file.unlink()
            try:
                file.parent.rmdir()
            except OSError:
                pass


def create_window(api=None, hidden=False):
    import webview
    # Browser drag regions must never claim buttons or other interactive children.
    webview.settings["DRAG_REGION_DIRECT_TARGET_ONLY"] = True
    if sys.platform == "win32":
        webview.settings["DRAG_REGION_SELECTOR"] = ".wenqi-browser-drag-disabled"
    api = api or DesktopApi()
    window = webview.create_window(APP_NAME, url=str(_resource("ui/index.html")), js_api=api,
        width=1180, height=820, min_size=(780, 620), frameless=True, easy_drag=False,
        shadow=True, hidden=hidden, background_color="#e9eef6", text_select=True)
    api._attach(window)

    def install_window_controls():
        if sys.platform == "win32":
            from scripts.window_controls import WindowsWindowControls
            controls = WindowsWindowControls(window)
            controls.install()
            api._window_controls = controls
            api._publish_window_state()

    window.events.shown += install_window_controls
    window.events.maximized += lambda: api._publish_window_state(maximized=True, minimized=False)
    window.events.minimized += lambda: api._publish_window_state(minimized=True)
    window.events.restored += lambda: api._publish_window_state(maximized=False, minimized=False)

    def bind_drop():
        from webview.dom import DOMEventHandler

        def drop(event):
            paths = [item.get("pywebviewFullPath") for item in event.get("dataTransfer", {}).get("files", [])]
            paths = [p for p in paths if p]
            if not paths:
                return
            try:
                state = api.add_files(paths)
                window.evaluate_js("window.wenqi.receiveState(" + json.dumps(state, ensure_ascii=False) + ")")
            except Exception as exc:
                window.evaluate_js("window.wenqi.toast(" + json.dumps(str(exc)) + ")")

        zone = window.dom.get_element("#drop-zone")
        if zone:
            zone.on("drop", DOMEventHandler(drop, prevent_default=True, stop_propagation=True))

    window.events.loaded += bind_drop
    def allow_close():
        if api._job["running"]:
            return False
        api._window_closing = True
        return True
    window.events.closing += allow_close
    window.events.closed += api._cleanup
    return window, api


def launch():
    import webview
    webview.settings["OPEN_DEVTOOLS_IN_DEBUG"] = False
    webview.settings["ALLOW_DOWNLOADS"] = False
    create_window()
    webview.start(gui="edgechromium" if sys.platform == "win32" else None,
                  icon=str(_resource("assets/icon.ico" if sys.platform == "win32" else "assets/icon.png")),
                  localization={"global.quitConfirmation": "确定退出文齐？"})
