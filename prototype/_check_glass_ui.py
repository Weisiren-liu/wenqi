"""Exercise and capture the real local browser desktop interface on Windows."""

import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import time
import traceback
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import docformat_gui as model
from scripts.desktop_ui import create_window
import webview
from _screenshot import capture


def main():
    destination = Path(__file__).parent / "_glass_shots"
    destination.mkdir(exist_ok=True)
    failure = []
    with TemporaryDirectory(prefix="wenqi-glass-check-") as work:
        model.CONFIG_FILE = Path(work) / "custom_settings.json"
        window, api = create_window()

        def run():
            try:
                window.events.loaded.wait(20)
                deadline = time.monotonic() + 20
                while time.monotonic() < deadline:
                    if window.evaluate_js("Boolean(window.wenqi && document.querySelector('#preset-select').options.length)"):
                        break
                    time.sleep(.25)
                else:
                    raise AssertionError("The local UI bridge did not initialize")
                time.sleep(1)
                print("BRIDGE READY", flush=True)
                from webview.dom import _dnd_state
                assert _dnd_state["num_listeners"] > 0, "Native file drop is not registered"
                print(window.evaluate_js("JSON.stringify({width:innerWidth,height:innerHeight,glass:getComputedStyle(document.querySelector('.document-surface')).backdropFilter})"), flush=True)
                hwnd = int(window.native.Handle.ToInt64())
                capture(hwnd, str(destination / "01-documents.png"))
                window.evaluate_js("window.wenqi.setView('format')")
                time.sleep(.7)
                capture(hwnd, str(destination / "02-format.png"))
                window.evaluate_js("document.querySelector('[data-key=body][data-field=font_cn]').click()")
                assert window.evaluate_js("document.querySelectorAll('#font-choices [role=option]').length") >= len(api.get_state()["fonts"])
                assert window.evaluate_js("document.querySelector('#font-search').value") == ""
                time.sleep(.3)
                capture(hwnd, str(destination / "10-font-picker.png"))
                window.evaluate_js("var search=document.querySelector('#font-search');search.value='宋体';search.dispatchEvent(new Event('input',{bubbles:true}));Array.from(document.querySelectorAll('#font-choices [role=option]')).find(option=>option.textContent==='宋体').click()")
                deadline = time.monotonic() + 5
                while api.get_state()["effective"]["body"]["font_cn"] != "宋体" and time.monotonic() < deadline:
                    time.sleep(.1)
                assert api.get_state()["effective"]["body"]["font_cn"] == "宋体"
                assert window.evaluate_js("document.querySelector('#font-picker').hidden")
                print("FONT PICKER: full list with a nonempty value, Chinese search and persisted selection: OK", flush=True)
                window.evaluate_js("document.querySelector('[data-key=body][data-field=align]').nextElementSibling.click()")
                assert window.evaluate_js("document.querySelectorAll('#select-picker [role=option]').length") == 4
                assert window.evaluate_js("getComputedStyle(document.querySelector('#select-picker')).backdropFilter.includes('url')")
                time.sleep(.3)
                capture(hwnd, str(destination / "13-select-picker.png"))
                window.evaluate_js("document.querySelector('#select-choice-2').click()")
                deadline = time.monotonic() + 5
                while api.get_state()["effective"]["body"]["align"] != "right" and time.monotonic() < deadline:
                    time.sleep(.1)
                assert api.get_state()["effective"]["body"]["align"] == "right"
                assert window.evaluate_js("document.querySelector('[data-key=body][data-field=align]').nextElementSibling.textContent") == "右对齐"
                print("ALIGNMENT PICKER: shared glass style, four choices and persisted selection: OK", flush=True)
                for category, setting, chosen, shot in (
                        ("page", "page_number_style", "plain", "16-page-style-picker.png"),
                        ("page", "page_number_position", "center", "17-page-position-picker.png"),
                        ("behavior", "space_handling", "keep_all", "18-space-picker.png")):
                    original = api.get_state()["effective"][setting]
                    selector = f'[data-key="{setting}"]'
                    window.evaluate_js("window.wenqi.setCategory(" + json.dumps(category) + ");document.querySelector(" + json.dumps(selector) + ").nextElementSibling.click()")
                    assert window.evaluate_js("!document.querySelector('#select-picker').hidden")
                    time.sleep(.2)
                    capture(hwnd, str(destination / shot))
                    window.evaluate_js("var select=document.querySelector(" + json.dumps(selector) + ");var index=Array.from(select.options).findIndex(option=>option.value===" + json.dumps(chosen) + ");document.querySelector('[data-select-index=\"'+index+'\"]').click()")
                    deadline = time.monotonic() + 5
                    while api.get_state()["effective"][setting] != chosen and time.monotonic() < deadline:
                        time.sleep(.1)
                    assert api.get_state()["effective"][setting] == chosen
                    window.evaluate_js("document.querySelector(" + json.dumps(selector) + ").nextElementSibling.click();var index=Array.from(document.querySelector(" + json.dumps(selector) + ").options).findIndex(option=>option.value===" + json.dumps(original) + ");document.querySelector('[data-select-index=\"'+index+'\"]').click()")
                    deadline = time.monotonic() + 5
                    while api.get_state()["effective"][setting] != original and time.monotonic() < deadline:
                        time.sleep(.1)
                    assert api.get_state()["effective"][setting] == original
                assert window.evaluate_js("Array.from(document.querySelectorAll('select')).every(select=>select.hidden)")
                print("ALL SELECTS: page style, page position and space handling share the picker and save correctly; no visible native menus: OK", flush=True)
                for category in ("page", "table", "behavior"):
                    window.evaluate_js("window.wenqi.setCategory(" + json.dumps(category) + ")")
                    time.sleep(.35)
                    assert window.evaluate_js("document.querySelector('#" + category + "-settings').getBoundingClientRect().height > 0")
                window.evaluate_js("window.wenqi.setCategory('paragraph');var input=document.querySelector('[data-key=body][data-field=size]');input.value=18;input.dispatchEvent(new Event('change',{bubbles:true}))")
                deadline = time.monotonic() + 5
                while api.get_state()["effective"]["body"]["size"] != 18 and time.monotonic() < deadline:
                    time.sleep(.1)
                assert api.get_state()["effective"]["body"]["size"] == 18, "Inline editing did not save"
                window.evaluate_js("window.wenqi.setView('documents');document.querySelector('#paste-text').click()")
                time.sleep(.5)
                capture(hwnd, str(destination / "03-paste.png"))
                window.evaluate_js("document.querySelector('#modal-close').click()")
                window.resize(800, 640)
                time.sleep(.8)
                capture(hwnd, str(destination / "04-narrow.png"))
                overflow = window.evaluate_js("JSON.stringify({body:document.body.scrollWidth>innerWidth,app:document.querySelector('#app').scrollWidth>innerWidth,button:document.querySelector('#start-processing').getBoundingClientRect().bottom>innerHeight})")
                print("NARROW", overflow, flush=True)
                assert not any(json.loads(overflow).values()), overflow
                window.resize(1180, 820)
                time.sleep(.5)
                source = Path(work) / "界面拖拽 测试.docx"
                document = Document()
                document.add_paragraph("工作安排")
                document.add_paragraph("这是正文, 请做好工作!")
                document.save(source)
                _dnd_state["paths"].append((source.name, str(source)))
                # Synthetic File objects have no OS handle. Preserve the real DOM
                # event / bridge route, but supply the native path cache ourselves.
                window.evaluate_js("(function(){const original=window.pywebview._jsApiCallback;window.pywebview._jsApiCallback=function(fn,params,id){if(params.event?.type==='drop'){window.pywebview._jsApiCallback=original;params={event:JSON.parse(window.pywebview.stringify(params.event)),nodeId:params.nodeId};}return original.call(this,fn,params,id);};return true})()")
                window.evaluate_js("(function(){const dt=new DataTransfer();dt.items.add(new File(['test']," + json.dumps(source.name) + "));document.querySelector('#drop-zone').dispatchEvent(new DragEvent('drop',{dataTransfer:dt,bubbles:true,cancelable:true}));return true})()")
                deadline = time.monotonic() + 5
                while not api.get_state()["files"] and time.monotonic() < deadline:
                    time.sleep(.1)
                assert api.get_state()["files"][0]["path"] == str(source), "File drop did not reach Python"
                time.sleep(.5)
                capture(hwnd, str(destination / "05-selected.png"))
                window.evaluate_js("document.querySelector('#start-processing').click()")
                deadline = time.monotonic() + 15
                while not api.get_state()["results"] and time.monotonic() < deadline:
                    time.sleep(.1)
                result = api.get_state()["results"][0]
                assert result["status"] == "success", result
                assert Document(result["path"]).paragraphs[1].runs[0].font.size.pt == 18
                time.sleep(.7)
                assert window.evaluate_js("!document.querySelector('#result-panel').hidden"), "Results did not render"
                capture(hwnd, str(destination / "06-result.png"))
                print("INLINE EDIT / DOM DROP / REAL DOCX PROCESSING / RESULT UI: OK", flush=True)
                print("SCREENSHOTS", str(destination.resolve()), flush=True)
            except Exception:
                failure.append(traceback.format_exc())
                print(failure[-1], flush=True)
            finally:
                window.destroy()

        webview.settings["OPEN_DEVTOOLS_IN_DEBUG"] = False
        webview.start(run, gui="edgechromium")
    if failure:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
