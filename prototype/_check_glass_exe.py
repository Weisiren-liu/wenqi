"""Verify the packaged Windows app through its real WebView2 and Python bridge.

Requires websockets. The debug port exists only in this isolated child process.
"""

import base64
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
from tempfile import TemporaryDirectory
import time
from urllib.request import urlopen

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from websockets.sync.client import connect
import win32com.client
import win32process
import win32gui
import win32api
import win32con
from _screenshot import capture, find_window

ROOT = Path(__file__).resolve().parents[1]


def main():
    exe = Path(sys.argv[1]).resolve() if len(sys.argv)>1 else ROOT / "dist" / "wenqi_windows.exe"
    shots = ROOT / "prototype" / "_glass_shots"
    shots.mkdir(exist_ok=True)
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        port = listener.getsockname()[1]
    with TemporaryDirectory(prefix="wenqi-packaged-ui-") as work:
        env = os.environ.copy()
        env["APPDATA"] = work
        env["WEBVIEW2_ADDITIONAL_BROWSER_ARGUMENTS"] = f"--remote-debugging-port={port}"
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        process = subprocess.Popen([str(exe)], env=env, cwd=ROOT, startupinfo=startup,
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        connection = None
        try:
            deadline = time.monotonic() + 30
            pages = []
            while time.monotonic() < deadline:
                assert process.poll() is None, f"EXE exited: {process.returncode}"
                try:
                    with urlopen(f"http://127.0.0.1:{port}/json/list", timeout=1) as response:
                        pages = [p for p in json.load(response) if p["type"] == "page"]
                    if pages and "index.html" in pages[0]["url"]:
                        break
                except OSError:
                    pass
                time.sleep(.2)
            assert pages, "WebView2 did not start"
            connection = connect(pages[0]["webSocketDebuggerUrl"], open_timeout=5, max_size=8*1024*1024)
            sequence = 0

            def command(method, **params):
                nonlocal sequence
                sequence += 1
                connection.send(json.dumps({"id": sequence, "method": method, "params": params}))
                deadline = time.monotonic() + 20
                while time.monotonic() < deadline:
                    response = json.loads(connection.recv(timeout=20))
                    if response.get("id") == sequence:
                        assert "error" not in response, response
                        return response.get("result", {})
                raise TimeoutError(method)

            def evaluate(expression):
                result = command("Runtime.evaluate", expression=expression,
                                 awaitPromise=True, returnByValue=True)
                assert "exceptionDetails" not in result, result
                return result.get("result", {}).get("value")

            def click(selector):
                point = evaluate("(()=>{const r=document.querySelector(" + json.dumps(selector) + ").getBoundingClientRect();return {x:r.x+r.width/2,y:r.y+r.height/2};})()")
                command("Input.dispatchMouseEvent", type="mousePressed", button="left", clickCount=1, **point)
                command("Input.dispatchMouseEvent", type="mouseReleased", button="left", clickCount=1, **point)

            def key(name, code):
                command("Input.dispatchKeyEvent", type="keyDown", key=name, windowsVirtualKeyCode=code)
                command("Input.dispatchKeyEvent", type="keyUp", key=name, windowsVirtualKeyCode=code)

            def wait_setting(setting, field, expected):
                deadline = time.monotonic() + 5
                while time.monotonic() < deadline:
                    data = evaluate("window.pywebview.api.get_state()")
                    value = data["effective"][setting]
                    if (value[field] if field else value) == expected:
                        return
                    time.sleep(.1)
                raise AssertionError(f"Setting did not save: {setting}.{field} = {expected}")

            def screenshot(name):
                # Native captures include the fixed window controls after scrolling;
                # CDP's compositor capture can omit untouched layers in WebView2.
                wmi = win32com.client.GetObject("winmgmts:")
                pids = {process.pid} | {int(item.ProcessId) for item in wmi.ExecQuery(
                    f"SELECT ProcessId FROM Win32_Process WHERE ParentProcessId={process.pid}")}
                for hwnd, _ in find_window("文齐"):
                    if win32process.GetWindowThreadProcessId(hwnd)[1] in pids:
                        capture(hwnd, str(shots / name))
                        return
                shot = command("Page.captureScreenshot", format="png", captureBeyondViewport=False)
                (shots / name).write_bytes(base64.b64decode(shot["data"]))

            deadline = time.monotonic() + 15
            while time.monotonic() < deadline:
                if evaluate("Boolean(window.pywebview?.api && document.querySelector('#preset-select').options.length)"):
                    break
                time.sleep(.2)
            else:
                raise AssertionError("The packaged Python / browser bridge did not initialize")
            assert evaluate("document.querySelectorAll('#element-table tbody tr').length") == 11
            assert evaluate("getComputedStyle(document.querySelector('.document-surface')).backdropFilter.includes('url')")
            assert evaluate("window.pywebview.api.window_state()")["nativeDrag"]
            wmi = win32com.client.GetObject("winmgmts:")
            pids = {process.pid} | {int(item.ProcessId) for item in wmi.ExecQuery(
                f"SELECT ProcessId FROM Win32_Process WHERE ParentProcessId={process.pid}")}
            hwnd = next(h for h, _ in find_window("文齐") if win32process.GetWindowThreadProcessId(h)[1] in pids)
            placement = win32gui.GetWindowRect(hwnd)
            workarea = win32api.GetMonitorInfo(win32api.MonitorFromWindow(hwnd, 2))["Work"]
            for _ in range(3):
                evaluate("window.pywebview.api.window_action('maximize').then(window.wenqi.receiveWindowState)")
                client = win32gui.GetClientRect(hwnd)
                assert (*win32gui.ClientToScreen(hwnd,client[:2]),*win32gui.ClientToScreen(hwnd,client[2:])) == workarea
                evaluate("window.pywebview.api.window_action('maximize').then(window.wenqi.receiveWindowState)")
                assert win32gui.GetWindowRect(hwnd) == placement
            evaluate("window.pywebview.api.window_action('minimize')")
            assert evaluate("window.pywebview.api.window_state()")["minimized"]
            win32gui.ShowWindow(hwnd,win32con.SW_RESTORE)
            time.sleep(.4)
            assert win32gui.GetWindowRect(hwnd) == placement
            # The managed WebView bounds can look correct while its HWND is displaced.
            child = win32gui.FindWindowEx(hwnd,0,None,None)
            assert win32gui.GetWindowRect(child) == placement, (win32gui.GetWindowRect(child),placement)
            print("PACKAGED WINDOW: work area, stable restore geometry, minimize/taskbar restore and aligned browser child: OK",flush=True)
            time.sleep(.5)
            screenshot("07-packaged-documents.png")
            preset_selector = '#preset-select + .select-trigger'
            original_preset = evaluate("window.pywebview.api.get_state()")["active"]
            for preset in ("official", original_preset):
                click(preset_selector)
                assert evaluate("!document.querySelector('#select-picker').hidden")
                if preset == "official":
                    time.sleep(.2)
                    screenshot("19-packaged-preset-picker.png")
                option_index = evaluate("Array.from(document.querySelector('#preset-select').options).findIndex(option=>option.value===" + json.dumps(preset) + ")")
                click(f"#select-choice-{option_index}")
                deadline = time.monotonic() + 5
                while evaluate("window.pywebview.api.get_state()")["active"] != preset and time.monotonic() < deadline:
                    time.sleep(.1)
                current = evaluate("window.pywebview.api.get_state()")
                assert current["active"] == preset
                assert evaluate("document.querySelector('#preset-select + .select-trigger').textContent") == current["effective"]["name"]
            evaluate("window.wenqi.setView('format');true")
            time.sleep(.4)
            screenshot("08-packaged-format.png")
            alignment = '[data-key="body"][data-field="align"] + .select-trigger'
            assert evaluate("document.querySelectorAll('#element-table .select-trigger').length") == 11
            assert evaluate("document.querySelectorAll('.select-trigger').length") == 15
            assert evaluate("Array.from(document.querySelectorAll('.select-control select')).every(select=>select.hidden)")
            click(alignment)
            assert evaluate("document.querySelectorAll('#select-picker [role=option]').length") == 4
            time.sleep(.2)
            assert evaluate("getComputedStyle(document.querySelector('#select-picker')).backdropFilter.includes('url')")
            assert evaluate("document.querySelector('#select-picker [aria-selected=true]').textContent") == "两端对齐"
            screenshot("14-packaged-select-picker.png")
            for index, (value, label) in enumerate((("left", "左对齐"), ("center", "居中"), ("right", "右对齐"), ("justify", "两端对齐"))):
                if index:
                    click(alignment)
                click(f"#select-choice-{index}")
                wait_setting("body", "align", value)
                assert evaluate("document.querySelector(" + json.dumps(alignment) + ").textContent") == label
                assert evaluate("document.querySelector('#select-picker').hidden")
            key("ArrowDown", 40)
            key("Home", 36)
            key("ArrowDown", 40)
            key("ArrowDown", 40)
            key("Enter", 13)
            wait_setting("body", "align", "right")
            click(alignment)
            key("Escape", 27)
            assert evaluate("document.querySelector('#select-picker').hidden")
            assert evaluate("document.activeElement.matches('.select-trigger')")
            click(alignment)
            key("Tab", 9)
            assert evaluate("document.querySelector('#select-picker').hidden")
            assert evaluate("document.activeElement.dataset.field") == "indent"
            click(alignment)
            click(".editor-heading h2")
            assert evaluate("document.querySelector('#select-picker').hidden")
            click(alignment)
            fonts = evaluate("window.pywebview.api.get_state()")["fonts"]
            default_english = evaluate("window.pywebview.api.get_state()")["effective"]["body"]["font_en"]
            chinese = '[data-key="body"][data-field="font_cn"]'
            english = '[data-key="body"][data-field="font_en"]'
            click(chinese)
            assert evaluate("document.querySelector('#select-picker').hidden"), "Font and alignment menus overlapped"
            assert evaluate("!document.querySelector('#font-picker').hidden")
            assert evaluate("document.querySelectorAll('#font-choices [role=option]').length") >= len(fonts)
            assert evaluate("document.querySelector('#font-search').value") == ""
            assert evaluate("document.activeElement.id") == "font-search"
            time.sleep(.2)
            screenshot("11-packaged-font-picker.png")
            command("Input.insertText", text="宋体")
            option_id = evaluate("Array.from(document.querySelectorAll('#font-choices [role=option]')).find(option=>option.textContent==='宋体').id")
            click("#" + option_id)
            wait_setting("body", "font_cn", "宋体")
            assert evaluate("document.querySelector('#font-picker').hidden")
            click(english + " + .font-toggle")
            command("Input.insertText", text="aRiAl")
            assert evaluate("document.querySelector('#font-choices [role=option]').textContent") == "Arial"
            key("ArrowDown", 40)
            key("ArrowUp", 38)
            key("Enter", 13)
            wait_setting("body", "font_en", "Arial")
            # Keyboard focus keeps manual entry and clearing a font available.
            evaluate("document.querySelector(" + json.dumps(english) + ").focus();document.querySelector(" + json.dumps(english) + ").select();true")
            command("Input.insertText", text="WenQi Custom Font")
            key("Tab", 9)
            wait_setting("body", "font_en", "WenQi Custom Font")
            evaluate("document.querySelector(" + json.dumps(english) + ").focus();document.querySelector(" + json.dumps(english) + ").select();true")
            key("Backspace", 8)
            key("Tab", 9)
            wait_setting("body", "font_en", default_english)
            assert "font_en" not in evaluate("window.pywebview.api.get_state()")["overrides"]["body"]
            click(chinese)
            key("Escape", 27)
            assert evaluate("document.querySelector('#font-picker').hidden")
            assert evaluate("document.activeElement.dataset.key==='body' && document.activeElement.dataset.field==='font_cn'")
            click(chinese)
            key("Tab", 9)
            assert evaluate("document.querySelector('#font-picker').hidden")
            assert evaluate("document.activeElement.dataset.field==='font_en'")
            # Check all font settings use the same full-list picker.
            for category, selector, setting, field in (
                    ("page", '[data-key="page_number_font"]', "page_number_font", None),
                    ("table", '[data-key="table"][data-field="font_cn"]', "table", "font_cn"),
                    ("table", '[data-key="table"][data-field="font_en"]', "table", "font_en")):
                evaluate("window.wenqi.setCategory(" + json.dumps(category) + ");true")
                click(selector)
                assert evaluate("document.querySelectorAll('#font-choices [role=option]').length") >= len(fonts)
                command("Input.insertText", text="Arial")
                key("Enter", 13)
                wait_setting(setting, field, "Arial")
            evaluate("window.wenqi.setCategory('paragraph');true")
            win32gui.SetWindowPos(hwnd,0,*placement[:2],800,640,win32con.SWP_NOZORDER|win32con.SWP_NOACTIVATE)
            time.sleep(.3)
            click('[data-key="title"][data-field="font_cn"]')
            assert evaluate("(()=>{const r=document.querySelector('#font-picker').getBoundingClientRect();return r.left>=0&&r.top>=0&&r.right<=innerWidth&&r.bottom<=innerHeight;})()"), "Font picker is clipped in a small window"
            time.sleep(.2)
            screenshot("12-packaged-font-narrow.png")
            click("#font-search")
            command("Input.insertText", text="WenQi Extra Font")
            assert evaluate("document.querySelector('#font-choices .custom-font').textContent.includes('WenQi Extra Font')")
            key("Escape", 27)
            click('[data-key="title"][data-field="font_cn"]')
            click('[data-category="page"]')
            assert evaluate("document.querySelector('#font-picker').hidden")
            evaluate("window.wenqi.setCategory('paragraph');true")
            click('[data-key="body"][data-field="align"] + .select-trigger')
            assert evaluate("(()=>{const r=document.querySelector('#select-picker').getBoundingClientRect();return r.left>=12&&r.top>=12&&r.right<=innerWidth-12&&r.bottom<=innerHeight-12;})()"), "Alignment picker is clipped in a small window"
            time.sleep(.2)
            screenshot("15-packaged-alignment-narrow.png")
            evaluate("document.querySelector('.table-scroller').scrollLeft=60;true")
            time.sleep(.1)
            assert evaluate("document.querySelector('#select-picker').hidden"), "Scrolling left an alignment menu detached from its field"
            win32gui.SetWindowPos(hwnd,0,*placement[:2],placement[2]-placement[0],placement[3]-placement[1],win32con.SWP_NOZORDER|win32con.SWP_NOACTIVATE)
            print("PACKAGED ALIGNMENT: all four choices, shared glass, save, keyboard, focus, dismiss, exclusive menus and small-window bounds: OK", flush=True)
            print(f"PACKAGED FONTS: {len(fonts)} fonts, full list, mouse selection, Chinese/case-insensitive search, keyboard, manual entry, restore, all settings and small-window bounds: OK", flush=True)
            for category, setting, chosen, count, shot in (
                    ("page", "page_number_style", "plain", 4, "20-packaged-page-style.png"),
                    ("page", "page_number_position", "center", 4, "21-packaged-page-position.png"),
                    ("behavior", "space_handling", "keep_all", 3, "22-packaged-space-handling.png")):
                evaluate("window.wenqi.setCategory(" + json.dumps(category) + ");true")
                selector = f'[data-key="{setting}"] + .select-trigger'
                original = evaluate("window.pywebview.api.get_state()")["effective"][setting]
                click(selector)
                assert evaluate("!document.querySelector('#select-picker').hidden")
                assert evaluate("document.querySelectorAll('#select-picker [role=option]').length") == count
                assert evaluate("(()=>{const r=document.querySelector('#select-picker').getBoundingClientRect();return r.left>=12&&r.top>=12&&r.right<=innerWidth-12&&r.bottom<=innerHeight-12;})()")
                time.sleep(.2)
                screenshot(shot)
                for value in (chosen, original):
                    if value == original:
                        click(selector)
                    index = evaluate("Array.from(document.querySelector('[data-key=" + setting + "]').options).findIndex(option=>option.value===" + json.dumps(value) + ")")
                    click(f"#select-choice-{index}")
                    wait_setting(setting, None, value)
                    assert evaluate("document.querySelector('#select-picker').hidden")
                    assert evaluate("document.querySelector('[data-key=" + setting + "]').selectedOptions[0].textContent===document.querySelector(" + json.dumps(selector) + ").textContent")
                click(selector)
                key("Escape", 27)
                assert evaluate("document.querySelector('#select-picker').hidden")
                assert evaluate("document.activeElement.matches(" + json.dumps(selector) + ")")
            assert evaluate("Array.from(document.querySelectorAll('select')).every(select=>select.hidden)")
            print("PACKAGED SELECTS: preset switch, page style, page position, space handling, shared menus, mouse selection, save, restore, focus and no visible native selects: OK", flush=True)
            for category in ("page", "table", "behavior"):
                evaluate("window.wenqi.setCategory(" + json.dumps(category) + ");true")
                assert evaluate("document.querySelector('#" + category + "-settings').getBoundingClientRect().height>0")
            evaluate("document.querySelector('#app-menu-toggle').click();document.querySelector('[data-action=about]').click();true")
            assert evaluate("document.querySelector('#modal-content').textContent.includes('KaguraNanaga')")
            evaluate("document.querySelector('#modal-close').click();true")
            source = Path(work) / "打包版处理 测试.docx"
            document = Document()
            document.add_paragraph("工作安排")
            document.add_paragraph("这是正文, 请做好工作!")
            document.save(source)
            evaluate("window.pywebview.api.add_files(" + json.dumps([str(source)]) + ").then(window.wenqi.receiveState)")
            evaluate("window.wenqi.setView('documents');document.querySelector('#start-processing').click();true")
            deadline = time.monotonic() + 15
            result = None
            while time.monotonic() < deadline:
                current = evaluate("window.pywebview.api.get_state()")
                if not current["job"]["running"] and current["results"]:
                    result = current["results"][0]
                    break
                time.sleep(.2)
            assert result and result["status"] == "success", result
            output = Document(result["path"])
            assert output.paragraphs[1].alignment == WD_ALIGN_PARAGRAPH.RIGHT
            assert output.paragraphs[1].text == "这是正文，请做好工作！"
            assert output.paragraphs[1].runs[0]._element.rPr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia") == "宋体"
            time.sleep(.7)
            assert evaluate("!document.querySelector('#result-panel').hidden")
            assert evaluate("document.querySelector('.open-result').getBoundingClientRect().bottom < document.querySelector('.action-dock').getBoundingClientRect().top")
            assert evaluate("document.querySelector('.titlebar').getBoundingClientRect().top >= 0"), "Window navigation scrolled away"
            screenshot("09-packaged-result.png")
            print("PACKAGED EXE: bridge, glass, 11 element rows, all settings tabs, attribution, real DOCX processing, visible result actions: OK", flush=True)
            evaluate("window.pywebview.api.window_action('close');true")
            process.wait(timeout=10)
        finally:
            if connection:
                connection.close()
            if process.poll() is None:
                # This helper owns the complete process tree; leave no test app behind.
                subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                process.wait(timeout=5)


if __name__ == "__main__":
    main()
