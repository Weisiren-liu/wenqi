#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""P3 端到端回归：真实 .docx 走一遍「智能一键处理」，对照 P1 之前的排版结果。

背景：P0/P1/P2 重写了配置路径、预设注册表和主页面（format_document 改走
preset= 完整 dict），此前所有验证用的都是合成配置，真实文档端到端没跑过。
本驱动用一份真实公文样例（标题 / 主送机关 / 多级标题 / 序号句 / 表格 /
落款 / 日期 / 附件，带混杂原始格式）做端到端对照：

  1. 基线：git worktree 检出 09d8aa5（P0/P1/P2 之前），驱动基线 GUI 的
     「智能一键处理」（步骤 1/2 修复标点 + 步骤 2/2 应用格式）
  2. 现行：当前 GUI 同一条智能流水线（自定义格式基准，无覆盖值）
  3. 两侧用同一份格式定义（当前 CUSTOM_PRESET_BASE），
     排版结果做语义指纹逐段对比：样式 / 对齐 / 缩进 / 行距 / 段前段后 /
     字体（中英）/ 字号 / 加粗 / 表格 / 页面设置 / 页眉页脚，必须逐项一致

用法：$env:PYTHONIOENCODING='utf-8'; python prototype/_check_p3_e2e.py
产物目录：prototype/_e2e_out/（gitignore，指纹文本留在里面供复盘）。
"""
import difflib
import json
import os
import shutil
import subprocess
import sys
import time
import traceback

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

BASELINE_COMMIT = '09d8aa5'          # P0/P1/P2 之前的基线
OUT = os.path.join(HERE, '_e2e_out')
BASELINE_WT = os.path.join(HERE, '_e2e_wt')   # 基线检出（保留复用，单独 gitignore）
INPUT = os.path.join(OUT, 'input', 'sample_notice.docx')
BASELINE_OUT = os.path.join(OUT, 'baseline_out.docx')
CURRENT_OUT = os.path.join(OUT, 'current_out.docx')


# ==================== 真实公文样例 ====================

def make_fixture(path):
    """一份带真实脏格式的公文样例：混杂字体、全角空格、中英数混排、表格。"""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()

    def para(text, font='仿宋_GB2312', size=16, bold=False, align=None,
             indent_chars=0, runs=None):
        p = doc.add_paragraph()
        chunks = runs if runs is not None else [(text, font, size, bold)]
        for chunk_text, chunk_font, chunk_size, chunk_bold in chunks:
            run = p.add_run(chunk_text)
            run.font.name = 'Times New Roman'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn('w:eastAsia'), chunk_font)
            run.font.size = Pt(chunk_size)
            run.bold = chunk_bold
        if align is not None:
            p.alignment = align
        if indent_chars:
            p.paragraph_format.first_line_indent = Pt(size * indent_chars)
        return p

    center = WD_ALIGN_PARAGRAPH.CENTER

    # 标题（故意用楷体 + 加粗，等着被排版纠正）
    para('关于开展2025年度安全生产大检查的通知', font='楷体_GB2312', size=22,
         bold=True, align=center)
    # 主送机关
    para('局机关各处室、各直属单位：', size=16)

    # 一、总体要求（含中英数混排、半角逗号脏数据）
    para('一、总体要求', font='黑体', size=16)
    para('　　为深入贯彻落实习近平总书记关于安全生产的重要论述，按照 GB/T 9704—2012 '
         '标准要求，结合我局实际,现就开展 2025 年度安全生产大检查工作通知如下。',
         size=16, indent_chars=2)
    para('（一）强化责任落实。各处室主要负责同志是第一责任人，要亲自研究、亲自部署、'
         '亲自督查，确保责任落实到岗到人。', size=16, indent_chars=2)
    para('（二）深化隐患排查。聚焦办公区域、档案库房、机房等重点部位，'
         '建立隐患台账，实行销号管理，确保不留死角、不留盲区。', size=16, indent_chars=2)

    para('1. 明确职责分工', font='楷体_GB2312', size=16)
    para('（1）办公室负责统筹协调、督办落实；', size=16, indent_chars=2)
    para('（2）各处室负责本领域自查自纠、立查立改；', size=16, indent_chars=2)

    # 序号句（触发「一是/二是」加粗）
    para('一是提高政治站位。要充分认识安全生产工作的极端重要性，坚决克服麻痹思想。'
         '二是压实工作责任。要坚持“谁主管、谁负责”，层层传导压力、级级压实责任。'
         '三是严格值班值守。要严格执行 24 小时值班和领导带班制度，确保信息畅通。',
         size=16, indent_chars=2)

    # 条款式段落（触发「第一条」加粗）
    para('　　第一条　各单位应当建立健全安全生产责任制，定期开展隐患排查治理，'
         '对发现的隐患实行闭环管理。', runs=[
             ('　　第一条　', '黑体', 16, False),
             ('各单位应当建立健全安全生产责任制，定期开展隐患排查治理，'
              '对发现的隐患实行闭环管理。', '宋体', 15, False)])

    # 二、时间安排
    para('二、时间安排', font='黑体', size=16)
    para('大检查自 2025 年12月3日起至 2025年12月31日结束，分动员部署、自查自纠、'
         '督导检查三个阶段进行，各阶段完成率分别不低于 35%、80%、100%。',
         size=16, indent_chars=2)

    # 表格：任务分解表（表头 + 数据行，数字列混排）
    caption = para('表1 安全生产大检查任务分解表', font='黑体', size=14, align=center)
    table = doc.add_table(rows=5, cols=4)
    table.style = 'Table Grid'
    data = [
        ('序号', '检查内容', '完成时限', '责任单位'),
        ('1', '办公区域用电安全检查', '12月10日', '办公室'),
        ('2', '档案库房消防设施检查', '12月15日', '档案处'),
        ('3', '机房设备运行检查', '12月20日', '信息中心'),
        ('4', '隐患整改回头看', '12月31日', '各处室'),
    ]
    for r, row_data in enumerate(data):
        for c, value in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = ''
            run = cell.paragraphs[0].add_run(value)
            run.font.name = 'Times New Roman'
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn('w:eastAsia'), '仿宋_GB2312')
            run.font.size = Pt(14)
            run.bold = (r == 0)

    # 三、工作要求
    para('三、工作要求', font='黑体', size=16)
    para('各处室要高度重视，认真组织，确保检查工作取得实效。对工作不力、失职渎职的，'
         '依纪依法严肃追究责任。检查中发现的重大问题，要第一时间报告。',
         size=16, indent_chars=2)

    # 结束语
    para('特此通知。', size=16, indent_chars=2)

    # 附件说明
    para('附件：1.安全生产大检查任务分解表', size=16)
    para('　　　2.隐患排查治理台账样式', size=16)

    # 落款 + 日期
    para('×××市市场监督管理局办公室', size=16, align=WD_ALIGN_PARAGRAPH.RIGHT)
    para('2025年12月3日', size=16, align=WD_ALIGN_PARAGRAPH.RIGHT)

    os.makedirs(os.path.dirname(path), exist_ok=True)
    doc.save(path)
    return path


# ==================== 排版结果语义指纹 ====================

def _len(value):
    if value is None:
        return None
    try:
        return round(value.pt, 2)
    except AttributeError:
        return round(float(value), 2)


def para_sig(p):
    pf = p.paragraph_format
    head = 'style=%s align=%s first=%s left=%s right=%s line=%s/%s before=%s after=%s' % (
        p.style.name if p.style else None, p.alignment,
        _len(pf.first_line_indent), _len(pf.left_indent), _len(pf.right_indent),
        pf.line_spacing, pf.line_spacing_rule,
        _len(pf.space_before), _len(pf.space_after))
    runs = []
    for r in p.runs:
        rPr = r._r.rPr
        fonts = rPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts') \
            if rPr is not None else None
        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
        ea = fonts.get(ns + 'eastAsia') if fonts is not None else None
        asc = fonts.get(ns + 'ascii') if fonts is not None else None
        runs.append('{%r ea=%s asc=%s sz=%s b=%s}' % (
            r.text, ea, asc, _len(r.font.size), r.bold))
    return head + ' :: ' + ' '.join(runs)


def fingerprint(path):
    from docx import Document
    doc = Document(path)
    out = []
    for i, sec in enumerate(doc.sections):
        out.append('section[%d] page=%sx%s margins=%s,%s,%s,%s hf=%s,%s' % (
            i, sec.page_width, sec.page_height,
            sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin,
            sec.header_distance, sec.footer_distance))
        for name, part in (('header', sec.header), ('footer', sec.footer)):
            for pi, p in enumerate(part.paragraphs):
                out.append('  %s.p[%d] %s' % (name, pi, para_sig(p)))
            out.append('  %s.xml %s' % (name, part.part.element.xml))
    out.append('--- body ---')
    for i, p in enumerate(doc.paragraphs):
        out.append('para[%d] %s' % (i, para_sig(p)))
    for ti, t in enumerate(doc.tables):
        out.append('table[%d] style=%s rows=%d cols=%d' % (
            ti, t.style.name if t.style else None, len(t.rows), len(t.columns)))
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    out.append('  cell[%d][%d].p[%d] %s' % (ri, ci, pi, para_sig(p)))
    return '\n'.join(out)


# ==================== 双侧执行 ====================

def ensure_worktree():
    wt = BASELINE_WT
    if not os.path.isdir(wt):
        # 目录可能被手动清理过，先清掉悬空注册
        subprocess.run(['git', 'worktree', 'prune'], cwd=ROOT, capture_output=True)
        proc = subprocess.run(['git', 'worktree', 'add', '--detach', wt, BASELINE_COMMIT],
                              cwd=ROOT, capture_output=True, text=True, encoding='utf-8')
        if proc.returncode != 0:
            raise RuntimeError('git worktree add failed: %s' % proc.stderr)
        print('  [基线检出]', wt, '@', BASELINE_COMMIT)
    return wt


def run_baseline(wt, settings_path, artifacts_path):
    runner = os.path.join(HERE, '_e2e_baseline_run.py')
    env = dict(os.environ, PYTHONIOENCODING='utf-8')
    proc = subprocess.run(
        [sys.executable, runner, wt, INPUT, BASELINE_OUT, artifacts_path, settings_path],
        cwd=ROOT, env=env, capture_output=True, text=True, encoding='utf-8')
    print('  [基线 stdout]', (proc.stdout or '').strip().splitlines()[-3:])
    if proc.returncode != 0:
        print('  [基线 stderr]', (proc.stderr or '')[-2000:])
    return proc.returncode == 0 and os.path.exists(BASELINE_OUT)


def run_current(artifacts_path):
    """当前 GUI 的「智能一键处理」：run_operation → smart 流水线。"""
    import tkinter as tk
    from tkinter import messagebox
    import docformat_gui as gui

    for name in ('showinfo', 'showwarning', 'showerror', 'askyesno',
                 'askyesnocancel', 'askretrycancel'):
        setattr(messagebox, name, lambda *a, **k: None)

    cfg = os.path.join(OUT, 'new_custom_settings.json')
    if os.path.exists(cfg):
        os.unlink(cfg)
    gui.CONFIG_FILE = gui.Path(cfg)
    gui.DocFormatApp._maybe_show_community_notice = lambda self: None

    root = tk.Tk()
    gui._configure_tk_high_dpi(root)
    app = gui.DocFormatApp(root)
    root.update()

    app.preset.set('custom')
    app.auto_open_var.set(False)
    app.input_files = [INPUT]
    app.output_file.set(CURRENT_OUT)
    app.operation.set('smart')

    with open(artifacts_path, 'w', encoding='utf-8') as f:
        json.dump({'effective_settings': app._get_selected_preset()},
                  f, ensure_ascii=False, indent=2)

    app.run_operation()

    # 真实 mainloop：工作线程经 after() 投递 Tk 调用，主线程必须在
    # mainloop 里（update() 轮询模式下会 RuntimeError: main thread
    # is not in main loop）
    state = {'deadline': time.time() + 300}

    def poll():
        if os.path.exists(CURRENT_OUT) and not app._processing_active:
            root.quit()
            return
        if time.time() > state['deadline']:
            root.quit()
            return
        root.after(100, poll)

    root.after(100, poll)
    root.mainloop()
    for _ in range(20):
        root.update()
        time.sleep(0.05)
    root.destroy()
    return os.path.exists(CURRENT_OUT)


# ==================== 主流程 ====================

def main():
    import docformat_gui as gui

    if os.path.isdir(OUT):
        shutil.rmtree(OUT)
    os.makedirs(OUT)

    failures = []

    def check(label, cond, detail=''):
        status = 'PASS' if cond else 'FAIL'
        print(f'[{status}] {label}' + (f' —— {detail}' if detail and not cond else ''))
        if not cond:
            failures.append(label)

    try:
        print('\n== 1. 生成真实公文样例 ==')
        make_fixture(INPUT)
        from docx import Document
        doc = Document(INPUT)
        check('样例含标题/正文/落款/表格',
              len(doc.tables) == 1 and len(doc.paragraphs) >= 18,
              f'paras={len(doc.paragraphs)} tables={len(doc.tables)}')
        check('样例表格含表头+数据行',
              len(doc.tables[0].rows) == 5 and len(doc.tables[0].columns) == 4)

        print('\n== 2. 双侧格式定义等值（同一份参数） ==')
        settings_path = os.path.join(OUT, 'settings.json')
        with open(settings_path, 'w', encoding='utf-8') as f:
            json.dump(gui.CUSTOM_PRESET_BASE, f, ensure_ascii=False, indent=2)

        wt = ensure_worktree()
        baseline_art = os.path.join(OUT, 'baseline_artifacts.json')
        current_art = os.path.join(OUT, 'current_artifacts.json')

        print('\n== 3. 基线（%s）智能一键处理 ==' % BASELINE_COMMIT)
        ok = run_baseline(wt, settings_path, baseline_art)
        check('基线处理完成并产出', ok)

        print('\n== 4. 当前 GUI 智能一键处理 ==')
        ok = run_current(current_art)
        check('当前处理完成并产出', ok)

        with open(baseline_art, encoding='utf-8') as f:
            baseline_used = json.load(f)['active_settings']
        with open(current_art, encoding='utf-8') as f:
            current_used = json.load(f)['effective_settings']
        meta = {'id', 'name', 'is_builtin', 'schema_version'}
        baseline_fmt = {k: v for k, v in baseline_used.items() if k not in meta}
        current_fmt = {k: v for k, v in current_used.items() if k not in meta}
        check('两侧实际使用的格式定义逐字段一致', baseline_fmt == current_fmt)
        if baseline_fmt != current_fmt:
            only_base = sorted(set(baseline_fmt) - set(current_fmt))
            only_curr = sorted(set(current_fmt) - set(baseline_fmt))
            differing = sorted(k for k in set(baseline_fmt) & set(current_fmt)
                               if baseline_fmt[k] != current_fmt[k])
            print('  [仅基线有]', only_base)
            print('  [仅当前有]', only_curr)
            print('  [取值不同]', differing)
            for k in differing[:5]:
                print('   ', k, 'baseline=', baseline_fmt[k], 'current=', current_fmt[k])

        print('\n== 5. 排版结果语义指纹对比 ==')
        fp_base = fingerprint(BASELINE_OUT)
        fp_curr = fingerprint(CURRENT_OUT)
        with open(os.path.join(OUT, 'fingerprint_baseline.txt'), 'w', encoding='utf-8') as f:
            f.write(fp_base)
        with open(os.path.join(OUT, 'fingerprint_current.txt'), 'w', encoding='utf-8') as f:
            f.write(fp_curr)

        base_lines = fp_base.splitlines()
        curr_lines = fp_curr.splitlines()
        check('段落/表格/页面数量一致',
              len(base_lines) == len(curr_lines),
              f'baseline={len(base_lines)} current={len(curr_lines)}')
        diff = list(difflib.unified_diff(base_lines, curr_lines,
                                         'baseline', 'current', lineterm='', n=0))
        check('排版结果逐项一致', not diff)
        if diff:
            print('  [差异前 40 行]')
            for line in diff[:40]:
                print('   ', line)

    except Exception:
        traceback.print_exc()
        failures.append('异常中断')

    print('\n========== 结果 ==========')
    if failures:
        print('失败 %d 项: %s' % (len(failures), '、'.join(failures)))
        return 1
    print('全部通过 ✓  产物目录:', OUT)
    return 0


if __name__ == '__main__':
    sys.exit(main())
